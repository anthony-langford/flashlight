"""CMHC Rental Market Survey workbooks (annual, October survey, purpose-built apartments).

Table 1.0 of each edition has one row per centre. We read the "Toronto CMA" row. Each edition
holds two October readings, so editions overlap and the newest edition wins on revisions.
Scope note: purpose-built rental apartment structures of three units and over. It excludes
condo apartments rented out, houses and basement apartments.
"""

from __future__ import annotations

import logging
import re
from datetime import date
from pathlib import Path

import httpx
import openpyxl

from ingest.models import Observation, SeriesDef

log = logging.getLogger(__name__)

BASE = (
    "https://assets.cmhc-schl.gc.ca/sites/cmhc/professional/housing-markets-data-and-research/"
    "housing-data-tables/rental-market/rental-market-report-data-tables"
)
# The file name changed in 2025. Both patterns are tried for every year.
FILE_NAMES = ["rmr-canada-{year}-en.xlsx", "rmr-{year}.xlsx"]
FIRST_EDITION = 2021  # earlier editions are not published at these URLs (checked 2026-10-14)
DEFAULT_DIR = Path(__file__).resolve().parent.parent.parent / "data" / "cmhc"

CENTRE = "Toronto CMA"
# Order of the numeric cells in a Table 1.0 row, after the centre name.
METRICS = [
    ("vacancy_rate", "prev"), ("vacancy_rate", "cur"),
    ("turnover_rate", "prev"), ("turnover_rate", "cur"),
    ("avg_rent_2br", "prev"), ("avg_rent_2br", "cur"),
    ("rent_change_2br", "prev"), ("rent_change_2br", "cur"),
]  # fmt: skip

_NUM = re.compile(r"^-?[\d,]+(\.\d+)?$")
_OCT = re.compile(r"Oct-(\d\d)")


def parse_workbook(path: Path) -> dict[str, list[Observation]]:
    """Return {metric: observations} for the Toronto CMA row of Table 1.0."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next((n for n in wb.sheetnames if "1.0" in n), None)
    if sheet is None:
        raise ValueError(f"{path.name}: no Table 1.0 sheet")
    rows = list(wb[sheet].iter_rows(values_only=True))
    cur_year = _current_year(rows, path)
    toronto = next((r for r in rows if r and r[0] == CENTRE), None)
    if toronto is None:
        raise ValueError(f"{path.name}: no {CENTRE} row")
    nums = [str(c).replace(",", "") for c in toronto[1:] if isinstance(c, str) and _NUM.match(c)]
    if len(nums) != len(METRICS):
        raise ValueError(f"{path.name}: expected {len(METRICS)} numbers, got {nums}")
    out: dict[str, list[Observation]] = {}
    for (metric, which), value in zip(METRICS, nums, strict=True):
        year = cur_year if which == "cur" else cur_year - 1
        out.setdefault(metric, []).append(
            Observation(obs_date=date(year, 10, 1), value=float(value))
        )
    return out


def _current_year(rows: list[tuple], path: Path) -> int:
    """Read the year from the 'Centre | Oct-24 | Oct-25 | ...' header and sanity check it."""
    for r in rows:
        if r and r[0] == "Centre":
            cells = [str(c) for c in r[1:] if c]  # merged cells leave empty gaps
            years = [int(m[1]) + 2000 for c in cells[:2] for m in [_OCT.search(c)] if m]
            if len(years) == 2 and years[1] == years[0] + 1:
                return years[1]
    raise ValueError(f"{path.name}: could not read the survey year from the header")


def download(client: httpx.Client, year: int, directory: Path = DEFAULT_DIR) -> Path | None:
    """Download one edition if not cached. Returns None when the edition does not exist."""
    for name in FILE_NAMES:
        path = directory / name.format(year=year)
        if path.exists():
            return path
    for name in FILE_NAMES:
        resp = client.get(f"{BASE}/{year}/{name.format(year=year)}", follow_redirects=True)
        if resp.status_code == 200 and "spreadsheet" in resp.headers.get("content-type", ""):
            directory.mkdir(parents=True, exist_ok=True)
            path = directory / name.format(year=year)
            path.write_bytes(resp.content)
            return path
    return None


def fetch(
    series: SeriesDef,
    client: httpx.Client,
    since: date | None = None,
    directory: Path = DEFAULT_DIR,
) -> list[Observation]:
    """`series.source_id` is one of the metric keys in METRICS. Editions are merged, newest wins."""
    merged: dict[date, Observation] = {}
    for year in range(FIRST_EDITION, date.today().year + 1):
        path = download(client, year, directory)
        if path is None:
            continue
        for obs in parse_workbook(path).get(series.source_id, []):
            merged[obs.obs_date] = obs  # later editions overwrite earlier readings
    if not merged:
        raise RuntimeError(f"no CMHC rental editions found for {series.source_id}")
    return [merged[d] for d in sorted(merged)]
